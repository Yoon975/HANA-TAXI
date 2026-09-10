"""
배차일지 제출용 내보내기.
- 관리용 xlsx는 그대로 두고, 제출용 파일에 하단 요약·결제란을 붙입니다.
- 제출용 xlsx는 항상 생성, HWP는 한글(한컴)이 설치된 PC에서만 시도합니다.
"""

from __future__ import annotations

import calendar
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from modules import dispatch_log
from modules.paths import CONFIG_DIR, DISPATCH_ROOT, ensure_storage

META_FILE = "dispatch_month_meta.json"
APPROVAL_TITLES = ["부장", "상무", "전무", "사장"]


def _meta_path() -> Path:
    ensure_storage()
    return CONFIG_DIR / META_FILE


def load_month_meta() -> dict[str, Any]:
    path = _meta_path()
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def save_month_meta(data: dict[str, Any]) -> Path:
    path = _meta_path()
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def month_key(year: int, month: int) -> str:
    return f"{year:04d}-{month:02d}"


def get_revenue_total(year: int, month: int) -> float:
    meta = load_month_meta().get(month_key(year, month)) or {}
    try:
        return float(meta.get("revenue_total") or 0)
    except (TypeError, ValueError):
        return 0.0


def set_revenue_total(year: int, month: int, total: float) -> None:
    data = load_month_meta()
    key = month_key(year, month)
    item = dict(data.get(key) or {})
    item["revenue_total"] = float(total)
    data[key] = item
    save_month_meta(data)


def add_revenue_total_from_agg(agg: pd.DataFrame) -> dict[str, float]:
    """수익 반영 시 월별 총액을 누적 저장(해당 월 합으로 덮어씀)."""
    if agg is None or agg.empty:
        return {}
    updated: dict[str, float] = {}
    work = agg.copy()
    work["_ym"] = work["business_date"].map(
        lambda d: month_key(int(d.year), int(d.month)) if hasattr(d, "year") else ""
    )
    for ym, g in work.groupby("_ym"):
        if not ym:
            continue
        total = float(g["revenue"].fillna(0).sum())
        y, m = int(ym[:4]), int(ym[5:7])
        set_revenue_total(y, m, total)
        updated[ym] = total
    return updated


def compute_summary(
    df: pd.DataFrame,
    year: int,
    month: int,
    *,
    revenue_total: float | None = None,
) -> dict[str, Any]:
    work = dispatch_log._normalize_df(df) if df is not None else pd.DataFrame()
    total_days = int(pd.to_numeric(work.get("근무일수"), errors="coerce").fillna(0).sum()) if len(work) else 0
    n_people = max(len(work), 1)
    per_person = total_days / n_people
    rev = float(revenue_total) if revenue_total is not None else get_revenue_total(year, month)
    avg = (rev / total_days) if total_days > 0 else 0.0
    return {
        "year": year,
        "month": month,
        "total_work_days": total_days,
        "per_person": per_person,
        "n_people": n_people,
        "revenue_total": rev,
        "avg_per_work_day": avg,
        "line1": f"■ 실 근무일수 : {total_days:,.3f} (1인기준 {per_person:,.3f})",
        "line2": (
            f"■ {month}월 운송수입금총액 : {rev:,.0f}  "
            f"실 근무일수 {total_days:,} 평균 {avg:,.0f}"
        ),
    }


def submission_xlsx_path(year: int, month: int) -> Path:
    ensure_storage()
    DISPATCH_ROOT.mkdir(parents=True, exist_ok=True)
    return DISPATCH_ROOT / f"배차일지_{year:04d}-{month:02d}_제출.xlsx"


def submission_hwp_path(year: int, month: int) -> Path:
    ensure_storage()
    DISPATCH_ROOT.mkdir(parents=True, exist_ok=True)
    return DISPATCH_ROOT / f"배차일지_{year:04d}-{month:02d}_제출.hwp"


def _thin() -> Border:
    side = Side(style="thin", color="000000")
    return Border(left=side, right=side, top=side, bottom=side)


def export_submission_xlsx(
    year: int,
    month: int,
    *,
    revenue_total: float | None = None,
    df: pd.DataFrame | None = None,
) -> tuple[Path, dict[str, Any]]:
    """사진과 같은 하단 요약·결제란이 포함된 제출용 xlsx."""
    if df is None:
        _, df = dispatch_log.load_month(year, month)
    else:
        df = dispatch_log._normalize_df(df)

    if revenue_total is not None:
        set_revenue_total(year, month, revenue_total)

    summary = compute_summary(df, year, month, revenue_total=revenue_total)
    path = submission_xlsx_path(year, month)
    last_day = calendar.monthrange(year, month)[1]
    title = f"배차일지 ({year}년 {month}월 1일 ~ {month}월 {last_day}일)"

    wb = Workbook()
    ws = wb.active
    ws.title = "배차일지"

    headers = dispatch_log.ALL_COLS
    ncols = len(headers)
    thin = _thin()
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c0 = ws.cell(1, 1, title)
    c0.font = Font(name="맑은 고딕", size=14, bold=True)
    c0.alignment = center

    for col, name in enumerate(headers, start=1):
        cell = ws.cell(2, col, name)
        cell.font = Font(name="맑은 고딕", size=9, bold=True)
        cell.alignment = center
        cell.border = thin

    for r_i, (_, row) in enumerate(df.iterrows(), start=3):
        for c_i, name in enumerate(headers, start=1):
            val = row.get(name, "")
            cell = ws.cell(r_i, c_i, "" if pd.isna(val) else val)
            cell.font = Font(name="맑은 고딕", size=9)
            cell.alignment = center
            cell.border = thin

    data_end = 2 + max(len(df), 0)
    sum_row = data_end + 2
    ws.merge_cells(start_row=sum_row, start_column=1, end_row=sum_row, end_column=ncols)
    s1 = ws.cell(sum_row, 1, summary["line1"])
    s1.font = Font(name="맑은 고딕", size=11, bold=True)
    s1.alignment = left

    ws.merge_cells(start_row=sum_row + 1, start_column=1, end_row=sum_row + 1, end_column=ncols)
    s2 = ws.cell(sum_row + 1, 1, summary["line2"])
    s2.font = Font(name="맑은 고딕", size=11, bold=True)
    s2.alignment = left

    label_row = sum_row + 3
    ws.cell(label_row, 1, "표 1").font = Font(name="맑은 고딕", size=10)

    # 우측 결제란: 결제 | 부장 상무 전무 사장
    stamp_row = sum_row + 2
    stamp_cols = 5
    start_col = max(ncols - stamp_cols + 1, 1)
    # 세로 "결제"
    ws.merge_cells(
        start_row=stamp_row,
        start_column=start_col,
        end_row=stamp_row + 2,
        end_column=start_col,
    )
    pay = ws.cell(stamp_row, start_col, "결\n제")
    pay.font = Font(name="맑은 고딕", size=11, bold=True)
    pay.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    pay.border = thin
    for r in range(stamp_row, stamp_row + 3):
        for c in range(start_col, start_col + stamp_cols):
            ws.cell(r, c).border = thin

    for i, title_name in enumerate(APPROVAL_TITLES):
        cell = ws.cell(stamp_row, start_col + 1 + i, title_name)
        cell.font = Font(name="맑은 고딕", size=10, bold=True)
        cell.alignment = center
        cell.border = thin
        # 도장칸
        ws.merge_cells(
            start_row=stamp_row + 1,
            start_column=start_col + 1 + i,
            end_row=stamp_row + 2,
            end_column=start_col + 1 + i,
        )
        stamp = ws.cell(stamp_row + 1, start_col + 1 + i, "")
        stamp.alignment = center
        stamp.border = thin

    # 열 너비
    widths = {"순서": 5, "차번": 8, "차종": 7, "성명": 8, "근무일수": 8}
    for col, name in enumerate(headers, start=1):
        if name in widths:
            ws.column_dimensions[get_column_letter(col)].width = widths[name]
        elif name.isdigit():
            ws.column_dimensions[get_column_letter(col)].width = 3.2
        else:
            ws.column_dimensions[get_column_letter(col)].width = 6

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[stamp_row + 1].height = 28
    ws.row_dimensions[stamp_row + 2].height = 28

    wb.save(path)
    return path, summary


def _hangul_available() -> bool:
    try:
        import win32com.client  # noqa: F401

        return True
    except ImportError:
        return False


def export_submission_hwp(
    year: int,
    month: int,
    *,
    revenue_total: float | None = None,
    df: pd.DataFrame | None = None,
) -> tuple[Path, dict[str, Any]]:
    """
    한글 COM으로 제출용 HWP 생성.
    표 + 요약 두 줄 + 결제란(부장·상무·전무·사장).
    """
    if df is None:
        _, df = dispatch_log.load_month(year, month)
    else:
        df = dispatch_log._normalize_df(df)

    if revenue_total is not None:
        set_revenue_total(year, month, revenue_total)

    summary = compute_summary(df, year, month, revenue_total=revenue_total)
    out_path = submission_hwp_path(year, month)

    try:
        import win32com.client as win32
    except ImportError as e:
        raise RuntimeError(
            "HWP 내보내기에는 pywin32가 필요합니다. pip install pywin32 후 다시 시도하세요."
        ) from e

    hwp = None
    try:
        try:
            hwp = win32.gencache.EnsureDispatch("HWPFrame.HwpObject")
        except Exception:
            hwp = win32.Dispatch("HWPFrame.HwpObject")
        try:
            hwp.RegisterModule("FilePathCheckDLL", "FilePathCheckerModule")
        except Exception:
            pass
        try:
            hwp.XHwpWindows.Item(0).Visible = False
        except Exception:
            pass

        last_day = calendar.monthrange(year, month)[1]
        title = f"배차일지 ({year}년 {month}월 1일 ~ {month}월 {last_day}일)"
        hwp.HAction.GetDefault("InsertText", hwp.HParameterSet.HInsertText.HSet)
        hwp.HParameterSet.HInsertText.Text = title + "\r\n"
        hwp.HAction.Execute("InsertText", hwp.HParameterSet.HInsertText.HSet)

        cols = len(dispatch_log.ALL_COLS)
        rows = len(df) + 1  # header + data
        hwp.HAction.GetDefault("TableCreate", hwp.HParameterSet.HTableCreation.HSet)
        hwp.HParameterSet.HTableCreation.Rows = int(rows)
        hwp.HParameterSet.HTableCreation.Cols = int(cols)
        hwp.HParameterSet.HTableCreation.WidthType = 2  # 쪽에 맞춤
        hwp.HParameterSet.HTableCreation.HeightType = 0
        hwp.HAction.Execute("TableCreate", hwp.HParameterSet.HTableCreation.HSet)

        def _put(text: str) -> None:
            hwp.HAction.GetDefault("InsertText", hwp.HParameterSet.HInsertText.HSet)
            hwp.HParameterSet.HInsertText.Text = str(text)
            hwp.HAction.Execute("InsertText", hwp.HParameterSet.HInsertText.HSet)

        def _next_cell() -> None:
            hwp.HAction.Run("TableRightCell")

        for i, name in enumerate(dispatch_log.ALL_COLS):
            _put(name)
            if i < cols - 1:
                _next_cell()

        for r_i in range(len(df)):
            hwp.HAction.Run("TableLowerCell")
            # 행 시작으로
            for _ in range(cols):
                hwp.HAction.Run("TableLeftCell")
            row = df.iloc[r_i]
            for c_i, col in enumerate(dispatch_log.ALL_COLS):
                val = row.get(col, "")
                if pd.isna(val):
                    val = ""
                _put(val)
                if c_i < cols - 1:
                    _next_cell()

        # 표 밖으로
        hwp.HAction.Run("TableBreakCell")
        _put("\r\n" + summary["line1"] + "\r\n")
        _put(summary["line2"] + "\r\n")
        _put("표 1\r\n\r\n")

        # 결제란 5열(결제+4직책) x 2행
        hwp.HAction.GetDefault("TableCreate", hwp.HParameterSet.HTableCreation.HSet)
        hwp.HParameterSet.HTableCreation.Rows = 2
        hwp.HParameterSet.HTableCreation.Cols = 5
        hwp.HParameterSet.HTableCreation.WidthType = 0
        hwp.HAction.Execute("TableCreate", hwp.HParameterSet.HTableCreation.HSet)
        titles = ["결제"] + APPROVAL_TITLES
        for i, t in enumerate(titles):
            _put(t)
            if i < 4:
                _next_cell()
        hwp.HAction.Run("TableLowerCell")
        for _ in range(5):
            hwp.HAction.Run("TableLeftCell")
        for i in range(5):
            _put("")
            if i < 4:
                _next_cell()

        abs_path = str(out_path.resolve())
        hwp.SaveAs(abs_path)
        return out_path, summary
    except Exception as e:
        raise RuntimeError(
            "한글(한컴오피스) HWP 내보내기에 실패했습니다. "
            "한글이 설치되어 있는지 확인하세요. "
            f"상세: {e}"
        ) from e
    finally:
        if hwp is not None:
            try:
                hwp.Quit()
            except Exception:
                pass


def export_submission(
    year: int,
    month: int,
    *,
    revenue_total: float | None = None,
    df: pd.DataFrame | None = None,
    try_hwp: bool = True,
) -> dict[str, Any]:
    """제출용 xlsx(+가능하면 HWP) 생성."""
    xlsx_path, summary = export_submission_xlsx(
        year, month, revenue_total=revenue_total, df=df
    )
    result: dict[str, Any] = {
        "xlsx": xlsx_path,
        "hwp": None,
        "summary": summary,
        "hwp_error": None,
    }
    if try_hwp:
        if not _hangul_available():
            result["hwp_error"] = "pywin32 없음 — HWP 생략, 제출용 xlsx만 생성"
        else:
            try:
                hwp_path, _ = export_submission_hwp(
                    year, month, revenue_total=revenue_total, df=df
                )
                result["hwp"] = hwp_path
            except Exception as e:
                result["hwp_error"] = str(e)
    return result
